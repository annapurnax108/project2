import datetime as dt

import openpyxl

from my_project.excel_fill import fill_template
from my_project.onboard import build_records


def make_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Onboarding"
    ws["A1"] = "guidelines"
    ws["Z10"] = "stale"
    wb.save(path)


def test_fill_writes_values_dates_and_extra_rows(tmp_path):
    src, dst = tmp_path / "t.xlsx", tmp_path / "out.xlsx"
    make_template(src)
    fixed = {"A": "Fixed Contact", "K": 110000147}
    people = [
        {"first_name": "Ann", "last_name": "One", "employment_status": "New Hire", "email": "a@x.com", "phone": "703-555-0100"},
        {"first_name": "Bo", "last_name": "Two", "employment_status": "Re-Hire", "rehire_details": "SAP 1", "email": "b@x.com", "phone": "703-555-0101"},
    ]
    records = build_records(people, fixed, dt.date(2026, 10, 7))
    fill_template(src, dst, records)

    ws = openpyxl.load_workbook(dst)["Onboarding"]
    assert ws["A1"].value == "guidelines"
    assert (ws["Z10"].value, ws["AA10"].value, ws["A10"].value, ws["K10"].value) == ("Ann", "One", "Fixed Contact", 110000147)
    assert (ws["Z11"].value, ws["AC11"].value, ws["AF11"].value) == ("Bo", "SAP 1", "b@x.com")
    serial = (dt.date(2026, 10, 7) - dt.date(1899, 12, 30)).days
    assert ws["C10"].value == serial
    assert ws["D10"].value == serial + 364
